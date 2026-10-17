"""Build renewal-model.xlsx (round 2, r2-economics-renewal).

Sheets
  README        what is formula vs value
  Inputs        mode (most-likely) value of every input, editable; source tag per row
  Calc          LIVE FORMULAS: the per-renewal chain for each line at the Inputs values (a deterministic
                "mode case"; it is NOT the Monte Carlo p50)
  MC_results    VALUES copied from renewal-results.csv (renewal_model.py, seed 20261017, n = 20,000)
  Pricing_grid  VALUES from renewal-pricing-grid.csv
  Buyers, MDE, Sensitivity   VALUES from the matching CSVs
After building, run LibreOffice headless to recalculate and compare Calc against mode_case() in Python.

Run: python3 build_xlsx.py && python3 build_xlsx.py --check <recalculated.xlsx>
"""
import csv
import os
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.dont_write_bytecode = True
sys.path.insert(0, HERE)
import renewal_model as M  # noqa: E402

LINES = list(M.LINES.keys())
COLS = {ln: get_column_letter(3 + i) for i, ln in enumerate(LINES)}  # C..F
HEAD = Font(bold=True)
INPUT_FILL = PatternFill("solid", fgColor="FFF4CC")

LINE_PARAMS = [
    ("P", "premium per renewal, INR"), ("c", "commission share of premium"), ("b0", "renewal rate with no calling"),
    ("s_pre", "share of baseline renewers who renew before the first call"), ("covH", "share of due list the human desk works"),
    ("persH", "P(saved | reached by human, would lapse)"), ("m_rpc", "minutes per right-party conversation"),
    ("r_future", "annual retention after a save (3-yr LTV)"), ("trough", "trough month / mean month due list"),
]
COMMON_PARAMS = list(M.COMMON.keys())
EXTRA = [("K_AI", 4, "AI attempts per contact (mode of design range 3-6)"),
         ("K_H", 2, "human attempts per worked contact [asm]"),
         ("cogs_launch", 5.88, "INR/min managed COGS at 100k min/month, MC p50 [l2v]"),
         ("cogs_scale", 3.06, "INR/min managed COGS at 1M min/month, MC p50 [l2]"),
         ("wa_price", M.WA_PRICE, "INR per WhatsApp utility template [prim, Meta rate card 2026-10-01]"),
         ("holdout", M.HOLDOUT, "share of due list held out [design]"),
         ("discount", M.DISCOUNT, "annual discount rate for LTV [asm]"),
         ("gm_target", 0.40, "VP gross-margin target for the fee floor"),
         ("talk_min_per_seat_month", M.HUMAN_MIN_PER_SEAT_MONTH, "l2v basis 200 talk-min/day x 24 days")]
FEES = {"2W_motor": 20, "4W_motor": 50, "health": 15, "life_yr2plus": 15}


def mode_case():
    """Deterministic evaluation of the same chain at mode values (mirrors the Calc sheet)."""
    out = {}
    C = {k: v[1] for k, v in M.COMMON.items()}
    K_AI, K_H = 4, 2
    for ln, L in M.LINES.items():
        x = {k: v[1] for k, v in L.items()}
        W = 1 - x["b0"]
        called = 1 - x["b0"] * x["s_pre"]
        share_wl = W / called
        m_mix = share_wl * x["m_rpc"] + (1 - share_wl) * x["m_rpc"] * C["short"]
        p_ai = min(max(C["p_ans"] * C["f1600"], 0.01), 0.95)
        a_ai = p_ai * C["q_rpc"]
        hit = 1 - (1 - a_ai) ** K_AI
        reach_ai = C["r_max"] * hit
        e_att = hit / a_ai
        att_ai = C["r_max"] * e_att + (1 - C["r_max"]) * K_AI
        answered = C["r_max"] * p_ai * e_att
        nonrpc = answered - reach_ai
        unans = att_ai - answered
        d_ai = W * reach_ai * min(1, x["persH"] * C["parity"])
        min_ai = called * (reach_ai * m_mix + nonrpc * C["m_nr"]) + d_ai * C["m_follow"]
        a_h = C["p_ans"] * C["q_rpc"]
        hit_h = 1 - (1 - a_h) ** K_H
        reach_h = C["r_max"] * hit_h
        answered_h = C["r_max"] * C["p_ans"] * hit_h / a_h
        d_h = W * x["covH"] * reach_h * x["persH"]
        min_h = called * x["covH"] * (reach_h * m_mix + (answered_h - reach_h) * C["m_nr"]) + d_h * C["m_follow"]
        hc = min_h * C["h_desk"]
        hcS = x["covH"] * C["seat_cost"] / C["due_per_seat"]
        v = x["c"] * x["P"]
        cost = {}
        for st, cg in (("launch", 5.88), ("scale", 3.06)):
            cost[st] = (min_ai * (cg + C["add1600"]) + called * unans * C["c_att"] + C["wa"] * M.WA_PRICE
                        + called * reach_ai * C["esc"] * C["m_esc"] * C["h"])
        fee = FEES[ln]
        rev = fee * (x["b0"] + d_ai)
        out[ln] = dict(d_ai_per_1000=1000 * d_ai, d_h_per_1000=1000 * d_h, min_ai_per_1000=1000 * min_ai,
                       vp_cost_launch_per_1000=1000 * cost["launch"], vp_cost_scale_per_1000=1000 * cost["scale"],
                       human_cost_T_per_1000=1000 * hc, human_cost_S_per_1000=1000 * hcS, value_per_save=v,
                       fc_max_nocall=d_ai * v / (x["b0"] + d_ai),
                       fc_max_vs_desk_T=((d_ai - d_h) * v + hc) / (x["b0"] + d_ai),
                       fc_floor_gm40_launch=cost["launch"] / (0.6 * (x["b0"] + d_ai)),
                       gm_launch_at_fee=(rev - cost["launch"]) / rev, gm_scale_at_fee=(rev - cost["scale"]) / rev,
                       yield_per_min_at_fee=rev / min_ai)
    return out


def build(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    notes = [
        "VP per-renewal unit economics (round 2, r2-economics-renewal, 2026-10-17)",
        "Inputs: mode values, editable (yellow). Calc: LIVE FORMULAS evaluating the chain at those values (a deterministic 'mode case').",
        "The mode case is NOT the Monte Carlo median. Decision numbers are in MC_results (values from renewal_model.py, seed 20261017, n=20,000).",
        "Pricing_grid, Buyers, MDE and Sensitivity are VALUES copied from the CSVs written by renewal_model.py; Expansion, TRAI_check and Incumbents from expansion_sizing.py; IRDAI_extracts from extract_irdai_handbook.py.",
        "Per due policy unless labelled 'per 1,000 due'. INR. FX not used (all inputs are INR).",
        "Assumption-heavy: AI/human outcome parity, persuasion, baseline renewal rates and desk productivity are NOT measured. The interviews must measure them.",
        "Source tags: [prim] primary, [comp] computed from IRDAI Handbook 2024-25, [sec] secondary, [vend] vendor claim, [asm] assumption.",
    ]
    for i, t in enumerate(notes, 1):
        ws.cell(row=i, column=1, value=t)
    ws.column_dimensions["A"].width = 140

    # Inputs
    wi = wb.create_sheet("Inputs")
    wi.append(["parameter", "description"] + LINES + ["low-high range (MC)", "source tag"])
    for c in wi[1]:
        c.font = HEAD
    pos = {}
    for key, desc in LINE_PARAMS:
        rng = "; ".join(f"{ln}: {M.LINES[ln][key][0]}-{M.LINES[ln][key][2]}" for ln in LINES)
        wi.append([key, desc] + [M.LINES[ln][key][1] for ln in LINES] + [rng, "see renewal_model.py comments"])
        pos[key] = wi.max_row
        for ln in LINES:
            wi[f"{COLS[ln]}{wi.max_row}"].fill = INPUT_FILL
    wi.append([])
    wi.append(["common parameter", "description", "value", "", "", "", "low-high range (MC)", "source tag"])
    for c in wi[wi.max_row]:
        c.font = HEAD
    cdesc = {
        "p_ans": "P(answered) per attempt [vend anchor Callyzer 46-48%]", "q_rpc": "P(right party | answered) [asm]",
        "r_max": "share reachable at all [asm]", "f1600": "pickup multiplier for AI on 1600xx [asm]",
        "parity": "AI persuasion / human persuasion [asm]", "m_nr": "minutes, answered wrong-party/not-now call [asm]",
        "short": "RPC minutes for would-renew-anyway / m_rpc [asm]", "wa": "WhatsApp utility messages per due policy [asm]",
        "c_att": "INR per unanswered attempt [asm]", "esc": "share of RPCs escalated to a human [asm]",
        "m_esc": "human minutes per escalation [asm]", "h": "INR per escalation minute [l2v 6-9]",
        "h_desk": "INR per desk talk-minute [asm on l2/l2v 6.3-8.9]", "m_follow": "extra minutes per saved customer [asm]",
        "seat_cost": "INR/month loaded desk seat [l3/l2]", "due_per_seat": "due renewals one seat works per month [asm; ASK]",
        "add1600": "INR/min 1600-series adder [asm; quotes pending]",
    }
    for key in COMMON_PARAMS:
        lo, mo, hi = M.COMMON[key]
        wi.append([key, cdesc.get(key, ""), mo, "", "", "", f"{lo}-{hi}", ""])
        pos[key] = wi.max_row
        wi[f"C{wi.max_row}"].fill = INPUT_FILL
    for key, val, desc in EXTRA:
        wi.append([key, desc, val])
        pos[key] = wi.max_row
        wi[f"C{wi.max_row}"].fill = INPUT_FILL
    wi.append([])
    wi.append(["fee_per_renewal", "fee per renewal collected (concept-test card)"] + [FEES[ln] for ln in LINES])
    pos["fee"] = wi.max_row
    for ln in LINES:
        wi[f"{COLS[ln]}{wi.max_row}"].fill = INPUT_FILL
    wi.column_dimensions["A"].width = 24
    wi.column_dimensions["B"].width = 58
    wi.column_dimensions["G"].width = 70

    def L(key, ln):   # line-specific input reference
        return f"Inputs!{COLS[ln]}{pos[key]}"

    def Cm(key):      # common input reference
        return f"Inputs!$C${pos[key]}"

    # Calc
    wc = wb.create_sheet("Calc")
    wc.append(["step", "formula meaning"] + LINES)
    for c in wc[1]:
        c.font = HEAD
    rows = {}

    def add(name, meaning, fn):
        wc.append([name, meaning] + [fn(ln) for ln in LINES])
        rows[name] = wc.max_row

    def R(name, ln):
        return f"{COLS[ln]}{rows[name]}"

    add("W", "would-lapse share = 1 - b0", lambda ln: f"=1-{L('b0', ln)}")
    add("called", "share unrenewed at first call = 1 - b0*s_pre", lambda ln: f"=1-{L('b0', ln)}*{L('s_pre', ln)}")
    add("share_wl", "would-lapse share of the called pool", lambda ln: f"={R('W', ln)}/{R('called', ln)}")
    add("m_mix", "RPC minutes, mixed pool", lambda ln: f"={R('share_wl', ln)}*{L('m_rpc', ln)}+(1-{R('share_wl', ln)})*{L('m_rpc', ln)}*{Cm('short')}")
    add("p_ai", "AI answer prob per attempt = p_ans*f1600 (clipped)", lambda ln: f"=MIN(0.95,MAX(0.01,{Cm('p_ans')}*{Cm('f1600')}))")
    add("a_ai", "AI RPC prob per attempt", lambda ln: f"={R('p_ai', ln)}*{Cm('q_rpc')}")
    add("hit_ai", "1-(1-a)^K_AI", lambda ln: f"=1-(1-{R('a_ai', ln)})^{Cm('K_AI')}")
    add("reach_ai", "AI right-party reach = r_max*hit", lambda ln: f"={Cm('r_max')}*{R('hit_ai', ln)}")
    add("eatt_ai", "expected attempts per reachable contact = hit/a", lambda ln: f"={R('hit_ai', ln)}/{R('a_ai', ln)}")
    add("att_ai", "attempts per called contact", lambda ln: f"={Cm('r_max')}*{R('eatt_ai', ln)}+(1-{Cm('r_max')})*{Cm('K_AI')}")
    add("answered_ai", "answered calls per called contact", lambda ln: f"={Cm('r_max')}*{R('p_ai', ln)}*{R('eatt_ai', ln)}")
    add("nonrpc_ai", "answered, not right party", lambda ln: f"={R('answered_ai', ln)}-{R('reach_ai', ln)}")
    add("unans_ai", "unanswered attempts per called contact", lambda ln: f"={R('att_ai', ln)}-{R('answered_ai', ln)}")
    add("d_ai", "incremental renewals per due vs no calling", lambda ln: f"={R('W', ln)}*{R('reach_ai', ln)}*MIN(1,{L('persH', ln)}*{Cm('parity')})")
    add("min_ai", "AI connected minutes per worked due", lambda ln: f"={R('called', ln)}*({R('reach_ai', ln)}*{R('m_mix', ln)}+{R('nonrpc_ai', ln)}*{Cm('m_nr')})+{R('d_ai', ln)}*{Cm('m_follow')}")
    add("a_h", "human RPC prob per attempt", lambda ln: f"={Cm('p_ans')}*{Cm('q_rpc')}")
    add("hit_h", "1-(1-a_h)^K_H", lambda ln: f"=1-(1-{R('a_h', ln)})^{Cm('K_H')}")
    add("reach_h", "human reach", lambda ln: f"={Cm('r_max')}*{R('hit_h', ln)}")
    add("answered_h", "human answered calls per worked contact", lambda ln: f"={Cm('r_max')}*{Cm('p_ans')}*{R('hit_h', ln)}/{R('a_h', ln)}")
    add("d_h", "human incremental renewals per due", lambda ln: f"={R('W', ln)}*{L('covH', ln)}*{R('reach_h', ln)}*{L('persH', ln)}")
    add("min_h", "human talk minutes per due", lambda ln: f"={R('called', ln)}*{L('covH', ln)}*({R('reach_h', ln)}*{R('m_mix', ln)}+({R('answered_h', ln)}-{R('reach_h', ln)})*{Cm('m_nr')})+{R('d_h', ln)}*{Cm('m_follow')}")
    add("human_cost_T", "desk cost per due, talk-minute method", lambda ln: f"={R('min_h', ln)}*{Cm('h_desk')}")
    add("human_cost_S", "desk cost per due, seat method", lambda ln: f"={L('covH', ln)}*{Cm('seat_cost')}/{Cm('due_per_seat')}")
    add("v", "commission per save, this year = c*P", lambda ln: f"={L('c', ln)}*{L('P', ln)}")
    add("v_ltv", "3-year commission value", lambda ln: f"={R('v', ln)}*(1+{L('r_future', ln)}/(1+{Cm('discount')})+({L('r_future', ln)}/(1+{Cm('discount')}))^2)")
    for st in ("launch", "scale"):
        add(f"vp_cost_{st}", f"VP cost per worked due at {st} COGS",
            lambda ln, st=st: f"={R('min_ai', ln)}*({Cm('cogs_' + st)}+{Cm('add1600')})+{R('called', ln)}*{R('unans_ai', ln)}*{Cm('c_att')}+{Cm('wa')}*{Cm('wa_price')}+{R('called', ln)}*{R('reach_ai', ln)}*{Cm('esc')}*{Cm('m_esc')}*{Cm('h')}")
    add("renewals_ai", "renewal rate on AI-worked list", lambda ln: f"={L('b0', ln)}+{R('d_ai', ln)}")
    add("fc_max_nocall", "buyer's max fee per renewal vs no calling", lambda ln: f"={R('d_ai', ln)}*{R('v', ln)}/{R('renewals_ai', ln)}")
    add("fc_max_vs_desk_T", "buyer's max fee per renewal vs own desk (talk method)", lambda ln: f"=(({R('d_ai', ln)}-{R('d_h', ln)})*{R('v', ln)}+{R('human_cost_T', ln)})/{R('renewals_ai', ln)}")
    add("fc_max_vs_desk_S", "buyer's max fee per renewal vs own desk (seat method)", lambda ln: f"=(({R('d_ai', ln)}-{R('d_h', ln)})*{R('v', ln)}+{R('human_cost_S', ln)})/{R('renewals_ai', ln)}")
    add("fc_floor_launch", "VP min fee per renewal for gm_target at launch", lambda ln: f"={R('vp_cost_launch', ln)}/((1-{Cm('gm_target')})*{R('renewals_ai', ln)})")
    add("fs_floor_launch", "VP min fee per incremental save for gm_target at launch", lambda ln: f"={R('vp_cost_launch', ln)}/((1-{Cm('gm_target')})*{R('d_ai', ln)})")
    add("revenue_per_due", "VP revenue per worked due at the fee", lambda ln: f"={L('fee', ln)}*{R('renewals_ai', ln)}")
    add("gm_launch", "VP gross margin at the fee, launch COGS", lambda ln: f"=({R('revenue_per_due', ln)}-{R('vp_cost_launch', ln)})/{R('revenue_per_due', ln)}")
    add("gm_scale", "VP gross margin at the fee, 1M min/month COGS", lambda ln: f"=({R('revenue_per_due', ln)}-{R('vp_cost_scale', ln)})/{R('revenue_per_due', ln)}")
    add("yield_per_min", "effective INR per AI-connected minute at the fee", lambda ln: f"={R('revenue_per_due', ln)}/{R('min_ai', ln)}")
    add("buyer_net_nocall_per_1000", "buyer net gain per 1,000 worked due vs no calling, INR", lambda ln: f"=1000*({R('d_ai', ln)}*{R('v', ln)}-{R('revenue_per_due', ln)})")
    add("kc1_due_per_month_3_seats_T", "due/month a 3-seat desk works (talk method)", lambda ln: f"=3*{Cm('talk_min_per_seat_month')}/{R('min_h', ln)}")
    add("kc1_due_per_month_3_seats_S", "due/month a 3-seat desk works (seat method)", lambda ln: f"=3*{Cm('due_per_seat')}/{L('covH', ln)}")
    wc.column_dimensions["A"].width = 30
    wc.column_dimensions["B"].width = 60
    for ln in LINES:
        wc.column_dimensions[COLS[ln]].width = 16
    wb._calc_rows = rows  # for the checker

    def copy_csv(name, fname, note):
        sh = wb.create_sheet(name)
        sh.append([note])
        sh["A1"].font = Font(italic=True)
        with open(os.path.join(HERE, fname), encoding="utf-8") as f:
            for i, r in enumerate(csv.reader(f)):
                vals = []
                for v in r:
                    try:
                        vals.append(float(v) if v not in ("", "True", "False") else v)
                    except ValueError:
                        vals.append(v)
                sh.append(vals)
                if i == 0:
                    for c in sh[sh.max_row]:
                        c.font = HEAD
        for col in range(1, 16):
            sh.column_dimensions[get_column_letter(col)].width = 18
        sh.column_dimensions["B"].width = 50
    copy_csv("MC_results", "renewal-results.csv", "VALUES from renewal_model.py (seed 20261017, n=20,000). Per 1,000 due where labelled.")
    copy_csv("Pricing_grid", "renewal-pricing-grid.csv", "VALUES from renewal_model.py: VP revenue, buyer gain and VP GM by fee and scheme.")
    copy_csv("Buyers", "renewal-buyers.csv", "VALUES: buyer archetypes at the concept-test fees (archetype sizes are assumptions).")
    copy_csv("MDE", "renewal-mde.csv", "VALUES: minimum detectable lift (pp) for a holdout, 80% power, alpha 0.05 two-sided.")
    copy_csv("Sensitivity", "renewal-sensitivity.csv", "VALUES: Spearman rank correlation of each input with outputs at the concept-test fee.")
    copy_csv("Expansion", "expansion-sizing.csv", "VALUES from expansion_sizing.py (seed 20261018): serviceable market by stage, INR crore/yr, as of 2026-10.")
    copy_csv("TRAI_check", "trai-reconciliation.csv", "VALUES: corrected direction of the TRAI 1600/140 comparison (l3v section 7).")
    copy_csv("Incumbents", "incumbent-reconciliation.csv", "VALUES: incumbent revenue (MCA via press = read-secondary) x assumed voice-agent share.")
    copy_csv("IRDAI_extracts", os.path.join("data", "irdai-handbook-2024-25-extracts.csv"), "VALUES: IRDAI Handbook 2024-25 tables T44, T50, T99, T103, T104 (extract_irdai_handbook.py).")
    wb.save(path)
    return rows


def check(recalc_path):
    rows = build(os.path.join(HERE, "_tmp_rows.xlsx"))
    os.remove(os.path.join(HERE, "_tmp_rows.xlsx"))
    wb = load_workbook(recalc_path, data_only=True)
    ws = wb["Calc"]
    mc = mode_case()
    pairs = {"d_ai": ("d_ai_per_1000", 1000), "d_h": ("d_h_per_1000", 1000), "min_ai": ("min_ai_per_1000", 1000),
             "vp_cost_launch": ("vp_cost_launch_per_1000", 1000), "vp_cost_scale": ("vp_cost_scale_per_1000", 1000),
             "human_cost_T": ("human_cost_T_per_1000", 1000), "human_cost_S": ("human_cost_S_per_1000", 1000),
             "v": ("value_per_save", 1), "fc_max_nocall": ("fc_max_nocall", 1), "fc_max_vs_desk_T": ("fc_max_vs_desk_T", 1),
             "fc_floor_launch": ("fc_floor_gm40_launch", 1), "gm_launch": ("gm_launch_at_fee", 1),
             "gm_scale": ("gm_scale_at_fee", 1), "yield_per_min": ("yield_per_min_at_fee", 1)}
    worst = 0.0
    for name, (key, mult) in pairs.items():
        for ln in LINES:
            xl = ws[f"{COLS[ln]}{rows[name]}"].value * mult
            py = mc[ln][key]
            rel = abs(xl - py) / max(abs(py), 1e-9)
            worst = max(worst, rel)
            if rel > 1e-6:
                print("MISMATCH", name, ln, xl, py)
    print("worst relative difference Calc vs Python mode case:", worst)
    for ln in LINES:
        print(ln, {k: round(v, 2) for k, v in mc[ln].items()})


if __name__ == "__main__":
    if len(sys.argv) > 2 and sys.argv[1] == "--check":
        check(sys.argv[2])
    else:
        build(os.path.join(HERE, "renewal-model.xlsx"))
        print("wrote renewal-model.xlsx")
