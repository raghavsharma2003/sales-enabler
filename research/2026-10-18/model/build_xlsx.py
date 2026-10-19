#!/usr/bin/env python3
"""Build pricing-model.xlsx from pricing_model.py (synthesis lane, 2026-10-19).

Sheets
  README       what is a live formula and what is a pasted value
  Inputs       every input with source, date, method and range (values; from pricing_model.INPUTS)
  Params       the editable deterministic case: FX, workload, telephony, infra, price, GST (named cells)
  Budget       LIVE FORMULAS: allowed AI cost per minute at the price, by GST treatment, margin and telephony
  Calc         LIVE FORMULAS: per cascade stack and tier x call length: STT, LLM, TTS, learning, AI, telephony,
               COGS, price needed at 40/50/60% margin, margin at the price. Column 'python_ai' holds the
               model's own value for the check
  Tiers        LIVE FORMULAS on Calc (3-minute row) + MC values from results/tiers.csv
  Live_g5      VALUES: Gemini Live paid meter (g5) and the calibrated model
  MC, Volume, Sensitivity, Latency, Calibration   VALUES from results/*.csv
Run:  python3 build_xlsx.py            (writes pricing-model.xlsx)
      python3 build_xlsx.py --check F  (F = a LibreOffice-recalculated copy; compares Calc to python)
"""
import csv
import os
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))
sys.dont_write_bytecode = True
sys.path.insert(0, HERE)
import pricing_model as M  # noqa: E402

RES = os.path.join(HERE, "results")
OUT = os.path.join(HERE, "pricing-model.xlsx")
BOLD = Font(bold=True)
INPUT_FILL = PatternFill("solid", fgColor="FFF4CC")
FX = 95.0

# per-stack row parameters (deterministic case); these mirror stt_cost / llm_cost / tts_cost in pricing_model.py
OR_FAC = 1.055 * 1.2
CPS = dict(az_neural=10.03, az_neural_c6=10.24, az_neural_c6g=9.66, az_hd=17.18, az_cnv=10.03)


def stt_params(kind):
    # mode, usd_per_sec, inr_per_sec, text_tok_per_exch, text_usd_per_tok, fixed_inr
    return {
        "az_full": ("full", 1.00 / 3600, 0, 0, 0, 0),
        "az_gated": ("gated", 1.00 / 3600, 0, 0, 0, 0),
        "gtl_full": ("full", 25 * 3.5e-6, 0, 14.6, 21e-6, 0),
        "gtl_gated": ("gated", 25 * 3.5e-6, 0, 14.6, 21e-6, 0),
        "sarvam_gated": ("gated", 0, 0.41 / 60, 0, 0, 0),
        "inworld_full": ("full", 0.10 / 3600, 0, 0, 0, 0),
        "selfhost": ("fixed", 0, 0, 0, 0, 0.0467),
    }[kind]


def llm_params(model, shape):
    pin, pc, pout, route = M.LLM_PRICE[model]
    sh = M.LLM_SHAPE[shape]
    cache = sh["cache"]
    if model in ("flashlite35", "gemma4_inworld"):
        cache = 0.0
    fac = OR_FAC if route == "openrouter" else 1.0
    return (pin, pc, pout, cache, sh["P0"], sh["g"], sh["waste"], sh["out"], fac)


def tts_params(kind, regime=2026):
    # usd_per_agent_s, inr_per_agent_s, lead_s, text_tok_per_exch, text_usd_per_tok, fixed_inr
    if kind in CPS:
        price = {"az_hd": 22e-6, "az_cnv": 24e-6}.get(kind, 15e-6)
        return (CPS[kind] * price, 0, 0, 0, 0, 0)
    if kind in ("gem_lite", "gem_flash"):
        pout = {("gem_lite", 2026): 6.0, ("gem_lite", 2027): 12.0, ("gem_flash", 2026): 9.0, ("gem_flash", 2027): 18.0}[(kind, regime)]
        pin = 0.5 if regime == 2026 else 1.0
        return (25 * pout / 1e6, 0, 0.25, 26, pin / 1e6, 0)
    if kind == "inworld_tts2f":
        return (14.67 * 7e-6, 0, 0, 0, 0, 0)
    if kind == "bulbul_v4f":
        return (0, 14.67 * 1.70 / 1000, 0, 0, 0, 0)
    if kind == "selfhost_veena":
        return (0, 0, 0, 0, 0, 0.139)
    raise KeyError(kind)


def calc_rows():
    rows = []
    for s in M.STACKS:
        if s["kind"] != "cascade":
            continue
        regimes = (2026, 2027) if s["tts"].startswith("gem") else (2026,)
        for regime in regimes:
            for L in M.LENGTHS:
                W = M.det_workload(L)
                py = float(M.stack_cost(s, W, {}, regime=regime)["ai"])
                rows.append(dict(id=s["id"] + ("@2027" if regime == 2027 else ""), name=s["name"], L=L, stt=stt_params(s["stt"]),
                                 llm=llm_params(*s["llm"]), s_extra=0, tts=tts_params(s["tts"], regime), learn=0, py=py))
    for t, d in M.TIER_DEFS.items():
        for L in M.LENGTHS:
            W = M.det_workload(L)
            py = float(M.tier_ai(d, W, {}, d["s_mode"])["ai"])
            rows.append(dict(id="tier:" + t, name=d["desc"], L=L, stt=stt_params(d["stt"]), llm=llm_params(*d["llm"]),
                             s_extra=d["s_mode"], tts=tts_params(d["tts"]), learn=1 if d["learn"] else 0, py=py))
    return rows


def read_csv(name):
    with open(os.path.join(RES, name)) as f:
        return list(csv.reader(f))


def paste(ws, rows, start_row=1):
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            if i > 0:
                try:
                    v = float(v)
                except (TypeError, ValueError):
                    pass
            ws.cell(row=start_row + i, column=j + 1, value=v)
        if i == 0:
            for j in range(len(r)):
                ws.cell(row=start_row, column=j + 1).font = BOLD


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    readme = [
        "VP voice-stack pricing model, 2026-10-19 (synthesis lane). Source of truth: model/pricing_model.py.",
        "Yellow cells on Params are editable inputs. Budget, Calc and Tiers are LIVE FORMULAS on Params.",
        "Live_g5, MC, Volume, Sensitivity, Latency and Calibration are VALUES pasted from results/*.csv.",
        "Calc covers cascades and tiers. Gemini Live is not a closed-form formula (per-invocation re-billing): use Live_g5.",
        "Deterministic case = g3 reference shape: 5 exchanges/min, agent 24 s/min, caller 24 s/min, FX 95, 2026 prices.",
        "OpenRouter-routed LLM rows include the 5.5% credit fee and a 1.2x ledger factor (g3 verify.md #7).",
        "Telephony per connected minute = rate x (L + 0.25)/L: 30-second pulses add 15 s per call on average.",
        "Prices are ex-GST. Input GST (incl. reverse-charge IGST on imported services) is assumed recoverable as ITC.",
        "Calc column 'diff' = formula AI minus python AI; it must be ~0 after recalculation (python3 build_xlsx.py --check).",
        "Every number on Inputs carries its lane file, date and method. No number here was invented: ranges marked 'assumption' are modelling choices.",
    ]
    for i, t in enumerate(readme, 1):
        ws.cell(row=i, column=1, value=t)
    ws.column_dimensions["A"].width = 140

    wi = wb.create_sheet("Inputs")
    cols = ["key", "value", "lo", "hi", "unit", "source", "date", "method", "note"]
    for j, c in enumerate(cols, 1):
        wi.cell(row=1, column=j, value=c).font = BOLD
    for i, r in enumerate(M.INPUTS, 2):
        for j, c in enumerate(cols, 1):
            wi.cell(row=i, column=j, value=r[c])
    for j, w in enumerate([28, 10, 8, 8, 60, 70, 11, 12, 70], 1):
        wi.column_dimensions[get_column_letter(j)].width = w

    wp = wb.create_sheet("Params")
    params = [("FX", 95.0, "INR per USD"), ("EXCH", 5.0, "exchanges per connected minute"), ("AGENT_S", 24.0, "agent speech s/min"),
              ("CALLER_S", 24.0, "caller speech s/min"), ("GR", 1.186, "gated STT seconds per caller second"),
              ("TEL", 0.43, "telephony INR per billed minute (0.38-0.48; 1600-series unpublished)"),
              ("INFRA", 0.03, "orchestration INR per connected minute (assumption 0.01-0.06)"),
              ("PRICE", 2.0, "price INR per connected minute"), ("GST_RATE", 0.18, "GST rate"),
              ("GST_INCL", 0, "1 if PRICE includes GST, 0 if GST is added on top")]
    wp.cell(row=1, column=1, value="name").font = BOLD
    wp.cell(row=1, column=2, value="value").font = BOLD
    wp.cell(row=1, column=3, value="meaning").font = BOLD
    for i, (n, v, m) in enumerate(params, 2):
        wp.cell(row=i, column=1, value=n)
        c = wp.cell(row=i, column=2, value=v)
        c.fill = INPUT_FILL
        wp.cell(row=i, column=3, value=m)
        wb.defined_names[n] = DefinedName(n, attr_text=f"Params!$B${i}")
    wp.column_dimensions["C"].width = 70

    wbud = wb.create_sheet("Budget")
    wbud["A1"] = "Allowed AI cost per minute = net price x (1 - margin) - telephony - infra (live formulas on Params)"
    wbud["A1"].font = BOLD
    hdr = ["GST treatment", "margin", "telephony 0.38", "0.43", "0.48", "1600 sens. 0.60", "1600 sens. 0.80", "at Params TEL"]
    for j, h in enumerate(hdr, 1):
        wbud.cell(row=3, column=j, value=h).font = BOLD
    r = 4
    for incl in (0, 1):
        for m in (0.4, 0.5, 0.6):
            wbud.cell(row=r, column=1, value="GST-inclusive" if incl else "ex-GST (GST on top)")
            wbud.cell(row=r, column=2, value=m)
            net = "PRICE/(1+GST_RATE)" if incl else "PRICE"
            for j, tel in enumerate((0.38, 0.43, 0.48, 0.60, 0.80), 3):
                wbud.cell(row=r, column=j, value=f"=ROUND({net}*(1-B{r})-{tel}-INFRA,3)")
            wbud.cell(row=r, column=8, value=f"=ROUND({net}*(1-B{r})-TEL-INFRA,3)")
            r += 1
    wbud.cell(row=r + 1, column=1, value="Price needed for a given AI cost: (AI + TEL + INFRA) / (1 - margin), ex-GST").font = BOLD
    wbud.cell(row=r + 2, column=1, value="AI cost INR/min")
    wbud.cell(row=r + 2, column=2, value=0.94).fill = INPUT_FILL
    for j, m in enumerate((0.4, 0.5, 0.6), 3):
        wbud.cell(row=r + 1, column=j, value=f"price at {int(m * 100)}%")
        wbud.cell(row=r + 2, column=j, value=f"=ROUND((B{r + 2}+TEL+INFRA)/(1-{m}),3)")
    wbud.column_dimensions["A"].width = 26

    wc = wb.create_sheet("Calc")
    head = ["id", "name", "L_min", "stt_mode", "stt_usd_per_s", "stt_inr_per_s", "stt_text_tok_per_exch", "stt_text_usd_per_tok", "stt_fixed_inr",
            "llm_in_usd_1m", "llm_cached_usd_1m", "llm_out_usd_1m", "cache_share", "P0_tok", "growth_tok_per_exch", "waste", "out_tok_per_exch",
            "llm_fee_x_ledger", "s_extra_tok", "tts_usd_per_agent_s", "tts_inr_per_agent_s", "tts_lead_s", "tts_text_tok_per_exch",
            "tts_text_usd_per_tok", "tts_fixed_inr", "learn",
            "STT", "LLM", "TTS", "LEARN", "AI", "TEL_per_min", "COGS", "price_at_40", "price_at_50", "price_at_60", "margin_at_PRICE",
            "python_ai", "diff"]
    for j, h in enumerate(head, 1):
        wc.cell(row=1, column=j, value=h).font = BOLD
    rows = calc_rows()
    for i, rr in enumerate(rows, 2):
        st, ll, tt = rr["stt"], rr["llm"], rr["tts"]
        vals = [rr["id"], rr["name"], rr["L"], st[0], st[1], st[2], st[3], st[4], st[5], ll[0], ll[1], ll[2], ll[3], ll[4], ll[5], ll[6], ll[7], ll[8],
                rr["s_extra"], tt[0], tt[1], tt[2], tt[3], tt[4], tt[5], rr["learn"]]
        for j, v in enumerate(vals, 1):
            wc.cell(row=i, column=j, value=v)
        f = {
            "STT": f'=IF(D{i}="full",60,IF(D{i}="gated",CALLER_S*GR,0))*(E{i}*FX+F{i})+EXCH*G{i}*H{i}*FX+I{i}',
            "LLM": f"=EXCH*R{i}*((N{i}+S{i}+O{i}*(MAX(EXCH*C{i},1)-1)/2)*P{i}*((1-M{i})*J{i}+M{i}*K{i})+Q{i}*L{i})/1000000*FX",
            "TTS": f"=(AGENT_S+V{i}*EXCH)*(T{i}*FX+U{i})+EXCH*W{i}*X{i}*FX+Y{i}",
            "LEARN": f"=IF(Z{i}=1,((214*C{i}+S{i}+1000)*0.44+400*1.76)/1000000*FX/C{i},0)",
            "AI": f"=AA{i}+AB{i}+AC{i}+AD{i}",
            "TEL_per_min": f"=TEL*(C{i}+0.25)/C{i}",
            "COGS": f"=AE{i}+AF{i}+INFRA",
            "price_at_40": f"=AG{i}/0.6", "price_at_50": f"=AG{i}/0.5", "price_at_60": f"=AG{i}/0.4",
            "margin_at_PRICE": f"=1-AG{i}/IF(GST_INCL=1,PRICE/(1+GST_RATE),PRICE)",
        }
        for j, h in enumerate(head, 1):
            if h in f:
                wc.cell(row=i, column=j, value=f[h])
        wc.cell(row=i, column=head.index("python_ai") + 1, value=round(rr["py"], 6))
        wc.cell(row=i, column=head.index("diff") + 1, value=f"=AE{i}-AL{i}")
    wc.column_dimensions["B"].width = 50
    wc.freeze_panes = "C2"
    calc_index = {(rr["id"], rr["L"]): i for i, rr in enumerate(rows, 2)}

    wt = wb.create_sheet("Tiers")
    th = ["tier", "stack", "price ex-GST (edit)", "price incl. GST", "AI 3-min (Calc)", "COGS 3-min (Calc)", "margin 3-min (formula)",
          "MC AI minute-weighted", "MC COGS minute-weighted", "MC margin minute-weighted", "MC P(call margin>=40%)", "price for 40% (MC)",
          "price for 50% (MC)", "dedicated-voice fee INR/month"]
    for j, h in enumerate(th, 1):
        wt.cell(row=1, column=j, value=h).font = BOLD
    mc = {r[0]: r for r in read_csv("tiers.csv")[1:]}
    mch = read_csv("tiers.csv")[0]
    for i, (t, d) in enumerate(M.TIER_DEFS.items(), 2):
        ci = calc_index[("tier:" + t, 3)]
        m = dict(zip(mch, mc[t]))
        wt.cell(row=i, column=1, value=t)
        wt.cell(row=i, column=2, value=d["desc"])
        wt.cell(row=i, column=3, value=d["price"]).fill = INPUT_FILL
        wt.cell(row=i, column=4, value=f"=ROUND(C{i}*(1+GST_RATE),2)")
        wt.cell(row=i, column=5, value=f"=Calc!AE{ci}")
        wt.cell(row=i, column=6, value=f"=Calc!AG{ci}")
        wt.cell(row=i, column=7, value=f"=1-F{i}/C{i}")
        for j, k in enumerate(["ai_inr_min_minute_weighted", "cogs_inr_min_minute_weighted", "gross_margin_minute_weighted",
                               "p_call_margin_ge_40", "price_for_40pct", "price_for_50pct", "voice_fee_inr_month"], 8):
            wt.cell(row=i, column=j, value=float(m[k]))
    wt.column_dimensions["B"].width = 70

    for name, fn in (("Live_g5", None), ("MC", "mc_by_stack_length.csv"), ("Volume", "volume_curve.csv"),
                     ("Sensitivity", "sensitivity_3min.csv"), ("Latency", "latency_by_stack.csv"),
                     ("Calibration", "live_calibration_vs_g5.csv"), ("Deterministic", "deterministic_by_stack_length.csv"),
                     ("Budget_values", "ai_budget_inr2.csv"), ("Tiers_by_length", "tiers_by_length_deterministic.csv")):
        w = wb.create_sheet(name)
        if fn is None:
            w.append(["Gemini Live AI cost, INR per connected minute, g5 paid meter (lanes/g5-gemini-paid/results/cost_per_minute_summary.csv), n=1-2 calls per cell"])
            w.append(["model", "compression", "1 min", "3 min", "5 min", "10 min", "10 min / 1 min"])
            for (mdl, comp), tab in M.G5.items():
                w.append([mdl, comp, tab[1], tab[3], tab[5], tab[10], round(tab[10] / tab[1], 2)])
            w.append([])
            w.append(["Tight compression (trigger 2500 / target 1500), 10 min: 3.1 2.85, 3.8 3.68; keeps ~17-27 s of audio history, so the agent forgets the call (g5)."])
            w.append(["Modelled at the reference workload and MC ranges: see MC sheet rows L31, L38, L38c. Model vs meter: see Calibration."])
            continue
        paste(w, read_csv(fn))
    wb.save(OUT)
    print("wrote", OUT, "with", len(rows), "Calc rows")


def check(path):
    wb = load_workbook(path, data_only=True)
    ws = wb["Calc"]
    head = [c.value for c in ws[1]]
    ia, ip, idf = head.index("AI"), head.index("python_ai"), head.index("diff")
    worst = 0.0
    n = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        if row[ia] is None:
            print("no cached value: recalc did not run")
            return 1
        worst = max(worst, abs(row[ia] - row[ip]))
        n += 1
    print(f"Calc rows checked: {n}; worst |formula - python| = {worst:.6f} INR/min")
    wb2 = wb["Budget"]
    print("Budget ex-GST 50% at TEL 0.43:", wb2["D5"].value, "| Tiers margins:", [wb["Tiers"].cell(row=i, column=7).value for i in range(2, 6)])
    return 0 if worst < 1e-6 else 2


if __name__ == "__main__":
    if len(sys.argv) > 2 and sys.argv[1] == "--check":
        sys.exit(check(sys.argv[2]))
    build()
